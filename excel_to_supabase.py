"""Convert the Excel tracker into INSERTs for the Supabase schema (test data / go-live import).

Usage:  pip install openpyxl
        python3 excel_to_supabase.py tracker.xlsx import.sql
Then run import.sql in Supabase SQL Editor, after supabase_schema.sql, on an empty database.

Rules applied (documented in the plan, section "Moving your Excel data"):
  * Supplier Master -> parties; C/O + Phone2 -> brokers
  * Purchase Tracker rows whose Invoice No starts with OPEN -> one OPENING bill (amount = Net Payable)
  * Other rows grouped by (supplier, invoice no, invoice date) -> one invoice with N lines
  * Same invoice no on a different date -> second one gets suffix -B (as you already did for J026-2026-27-2)
  * Missing payment term -> party default term
  * Weights rounded to 3 decimals
  * Payment log: rows without an amount are skipped; text dates (DD-MM-YYYY) and real dates both parsed;
    the Invoice No column is kept only as a remark – FIFO decides allocation now.
"""
import datetime as dt
import sys
from collections import OrderedDict, defaultdict
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "tracker.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data.sql"

wb = openpyxl.load_workbook(SRC, data_only=True)
EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def q(v):
    if v is None:
        return "null"
    return "'" + str(v).replace("'", "''") + "'"


def d3(x):
    return Decimal(str(x)).quantize(Decimal("0.001"), ROUND_HALF_UP)


def d2(x):
    return Decimal(str(x)).quantize(Decimal("0.01"), ROUND_HALF_UP)


def as_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Unparseable date: {v!r}")


def as_days(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):          # 90 formatted as a date in Excel -> 1900-03-30
        return (v - EXCEL_EPOCH).days
    return int(v)


def norm(s):
    return " ".join(str(s).split()).upper() if s is not None else None


out = []
notes = []

# ---------------------------------------------------------------- parties + brokers
sm = wb["Supplier Master"]
parties = OrderedDict()
brokers = OrderedDict()
for r in range(2, sm.max_row + 1):
    name = sm.cell(r, 1).value
    if not name:
        continue
    name = " ".join(str(name).split())
    co = sm.cell(r, 9).value
    co = " ".join(str(co).split()).title() if co else None
    phone2 = sm.cell(r, 10).value
    if co:
        brokers[co] = str(phone2) if phone2 else None
    parties[norm(name)] = dict(
        name=name,
        contact=(str(sm.cell(r, 2).value).strip() if sm.cell(r, 2).value else None),
        phone=(str(sm.cell(r, 3).value) if sm.cell(r, 3).value else None),
        email=sm.cell(r, 4).value,
        address=sm.cell(r, 5).value,
        term=as_days(sm.cell(r, 6).value) or 90,
        remarks=sm.cell(r, 8).value,
        broker=co,
    )

for b, ph in brokers.items():
    out.append(f"insert into public.brokers (name, phone) values ({q(b)}, {q(ph)});")
for p in parties.values():
    broker_sql = f"(select id from public.brokers where name = {q(p['broker'])})" if p["broker"] else "null"
    out.append(
        "insert into public.parties (name, contact_person, phone, email, address, default_term_days, remarks, broker_id) "
        f"values ({q(p['name'])}, {q(p['contact'])}, {q(p['phone'])}, {q(p['email'])}, {q(p['address'])}, "
        f"{p['term']}, {q(p['remarks'])}, {broker_sql});"
    )

# ---------------------------------------------------------------- invoices
pt = wb["Purchase Tracker"]
invoices = OrderedDict()   # key -> header
seen_no = defaultdict(set)  # (party, invoice_no) -> dates
for r in range(4, pt.max_row + 1):
    inv_no, sup, idate, wt, rate = (pt.cell(r, c).value for c in range(1, 6))
    if inv_no in (None, "") and sup in (None, ""):
        continue
    inv_no = str(inv_no).strip()
    pkey = norm(sup)
    if pkey not in parties:
        raise SystemExit(f"Row {r}: supplier {sup!r} not in Supplier Master")
    idate = as_date(idate)
    term = as_days(pt.cell(r, 12).value)
    if term is None:
        term = parties[pkey]["term"]
        notes.append(f"Row {r} {inv_no}: no payment term -> using party default {term}")
    qc = pt.cell(r, 7).value or 0
    net_payable = pt.cell(r, 11).value

    if inv_no.upper().startswith("OPEN"):
        key = (pkey, inv_no, idate)
        invoices[key] = dict(party=pkey, no=inv_no, date=idate, term=term, type="OPENING",
                             opening=d2(net_payable), lines=[], row=r)
        continue

    dates = seen_no[(pkey, inv_no)]
    final_no = inv_no
    if dates and idate not in dates:
        final_no = inv_no + "-B"
        notes.append(f"Row {r}: {inv_no} already used on another date -> imported as {final_no}")
    dates.add(idate)
    key = (pkey, final_no, idate)
    if key not in invoices:
        invoices[key] = dict(party=pkey, no=final_no, date=idate, term=term, type="PURCHASE",
                             opening=None, lines=[], row=r)
    elif invoices[key]["term"] != term:
        notes.append(f"Row {r}: {final_no} line has a different term ({term}) – header keeps {invoices[key]['term']}")
    invoices[key]["lines"].append(dict(weight=d3(wt), rate=d2(rate), qc=d3(qc), row=r,
                                       excel_net=Decimal(str(net_payable))))
    if Decimal(str(qc)) != d3(qc):
        notes.append(f"Row {r} {final_no}: QC weight {qc} rounded to {d3(qc)}")

for inv in invoices.values():
    pname = parties[inv["party"]]["name"]
    if inv["type"] == "OPENING":
        out.append(
            "insert into public.purchase_invoices (party_id, bill_type, invoice_no, invoice_date, term_days, opening_amount, remarks) "
            f"values ((select id from public.parties where name = {q(pname)}), 'OPENING', {q(inv['no'])}, '{inv['date']}', "
            f"{inv['term']}, {inv['opening']}, {q('Excel row ' + str(inv['row']))});"
        )
    else:
        # Recompute exactly like the database does, then put any difference vs the Excel
        # Net Payable into round_off so the imported bill equals the Excel bill to the paisa.
        app_total = Decimal("0")
        for l in inv["lines"]:
            taxable = d2(l["weight"] * l["rate"]) - d2(l["qc"] * l["rate"])
            app_total += taxable + d2(taxable * Decimal("1.5") / 100)
        excel_total = d2(sum(l["excel_net"] for l in inv["lines"]))
        round_off = excel_total - app_total
        if abs(round_off) >= 1:
            notes.append(f"{inv['no']} ({pname}): Excel {excel_total} vs recomputed {app_total} "
                         f"-> round_off {round_off} (Excel weight was back-calculated from the amount)")
        out.append(
            "with h as (insert into public.purchase_invoices (party_id, invoice_no, invoice_date, term_days, tax_rate_percent, round_off, remarks) "
            f"values ((select id from public.parties where name = {q(pname)}), {q(inv['no'])}, '{inv['date']}', {inv['term']}, 1.5, {round_off}, "
            f"{q('Excel row(s) ' + ','.join(str(l['row']) for l in inv['lines']))}) returning id) "
            "insert into public.purchase_invoice_lines (invoice_id, line_no, unit, weight, rate, qc_return_weight) "
            + "select h.id, v.* from h, (values "
            + ", ".join(f"({i + 1}, 'ct', {l['weight']}, {l['rate']}, {l['qc']})" for i, l in enumerate(inv["lines"]))
            + ") as v(line_no, unit, weight, rate, qc);"
        )

# ---------------------------------------------------------------- payments
pl = wb["Partial Payment Log"]
skipped = 0
for r in range(4, pl.max_row + 1):
    pdate, sup, inv_no, _, amount, mode = (pl.cell(r, c).value for c in range(1, 7))
    if sup in (None, "") and amount in (None, ""):
        continue
    if amount in (None, "", 0):
        skipped += 1
        notes.append(f"Payment row {r} ({sup}, {pdate}): no amount -> skipped")
        continue
    pkey = norm(sup)
    pname = parties[pkey]["name"]
    pdate = as_date(pdate)
    mode = (mode or "RTGS/NEFT").strip()
    out.append(
        "insert into public.party_payments (party_id, payment_date, amount, mode_id, remarks) values ("
        f"(select id from public.parties where name = {q(pname)}), '{pdate}', {d2(amount)}, "
        f"(select id from public.payment_modes where name = {q(mode)}), "
        f"{q(f'Excel row {r}; was logged against {inv_no}')});"
    )

with open(OUT, "w") as f:
    f.write("-- Generated from " + SRC + "\nbegin;\n" + "\n".join(out) + "\ncommit;\n")
print(f"parties={len(parties)} brokers={len(brokers)} invoices={len(invoices)} "
      f"lines={sum(len(i['lines']) for i in invoices.values())} payments={sum(1 for o in out if 'party_payments' in o)}")
print("\n".join(notes))
